"""Phase 2 document extraction: DOCX and image OCR, with explicit uncertainty.

OCR is optional at runtime and uses the installed Tesseract binary via pytesseract.
No output from this module is written to an ERP.
"""
from __future__ import annotations

import hashlib
import io
from pathlib import Path
from typing import Any

from packages.ingestion.phase1 import IngestedDocument, IngestionError, MAX_FILE_BYTES, _safe_name

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
DOCX_EXTENSIONS = {".docx"}


def ingest_phase2_bytes(name: str, data: bytes) -> IngestedDocument:
    """Extract a DOCX or image into traceable text without submitting business data."""
    if not isinstance(data, bytes):
        raise IngestionError("Input must be bytes.")
    if not data:
        raise IngestionError("Input is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise IngestionError(f"Input exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB limit.")

    safe_name = _safe_name(name)
    suffix = Path(safe_name).suffix.lower()
    digest = hashlib.sha256(data).hexdigest()

    if suffix in DOCX_EXTENSIONS:
        try:
            from docx import Document
        except ImportError as exc:
            raise IngestionError("DOCX support requires python-docx. Install the FDIA dependencies.") from exc
        try:
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table_index, table in enumerate(document.tables, start=1):
                parts.append(f"[Table {table_index}]")
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts).strip()
            if not text:
                return IngestedDocument(
                    safe_name, "docx", digest, text="",
                    warnings=["DOCX contained no readable paragraph or table text."],
                )
            return IngestedDocument(safe_name, "docx", digest, text=text)
        except Exception as exc:
            raise IngestionError(f"DOCX parsing failed: {exc}") from exc

    if suffix in IMAGE_EXTENSIONS:
        try:
            from PIL import Image, ImageOps
        except ImportError as exc:
            raise IngestionError("Image OCR requires Pillow. Install the FDIA dependencies.") from exc
        try:
            image = Image.open(io.BytesIO(data))
            image.verify()
            image = Image.open(io.BytesIO(data))
            image = ImageOps.exif_transpose(image).convert("RGB")
        except Exception as exc:
            raise IngestionError(f"Image could not be decoded: {exc}") from exc

        try:
            import pytesseract
            from pytesseract import TesseractNotFoundError
        except ImportError as exc:
            raise IngestionError(
                "Image OCR requires pytesseract and the Tesseract OCR system binary."
            ) from exc

        try:
            extracted = pytesseract.image_to_data(
                image, output_type=pytesseract.Output.DICT, config="--psm 6"
            )
        except TesseractNotFoundError as exc:
            raise IngestionError(
                "Tesseract OCR binary was not found. Install Tesseract and ensure it is on PATH."
            ) from exc
        except Exception as exc:
            raise IngestionError(f"Image OCR failed: {exc}") from exc

        words: list[str] = []
        confidences: list[float] = []
        for word, raw_conf in zip(extracted.get("text", []), extracted.get("conf", [])):
            word = str(word).strip()
            if not word:
                continue
            words.append(word)
            try:
                confidence = float(raw_conf)
                if confidence >= 0:
                    confidences.append(confidence / 100.0)
            except (TypeError, ValueError):
                continue

        text = " ".join(words).strip()
        average_confidence = sum(confidences) / len(confidences) if confidences else 0.0
        warnings: list[str] = []
        if not text:
            warnings.append("OCR returned no readable text; manual review is required.")
        if average_confidence < 0.80:
            warnings.append(
                f"OCR confidence is low ({average_confidence:.0%}); verify extracted values manually."
            )
        return IngestedDocument(
            safe_name, "image_ocr", digest, text=text, warnings=warnings,
            confidence=average_confidence,
        )

    raise IngestionError(
        f"Phase 2 supports DOCX and image files ({', '.join(sorted(IMAGE_EXTENSIONS))}); "
        "use ingest_bytes for Phase 1 formats."
    )


def ingest_any_bytes(name: str, data: bytes, *, pasted_text: bool = False) -> IngestedDocument:
    """Unified entry point for Phase 1 formats plus Phase 2 DOCX and image OCR."""
    suffix = Path(_safe_name(name)).suffix.lower()
    if pasted_text or suffix not in DOCX_EXTENSIONS | IMAGE_EXTENSIONS:
        from packages.ingestion.phase1 import ingest_bytes
        return ingest_bytes(name, data, pasted_text=pasted_text)
    return ingest_phase2_bytes(name, data)
