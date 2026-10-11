"""Safe, deterministic Phase 1 ingestion primitives.

This module deliberately does not write to any ERP. It extracts text/table rows,
normalizes candidate records, validates deterministic rules, and emits dry-run previews.
PDF extraction uses pypdf when installed; image OCR and browser automation are later phases.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


class IngestionError(ValueError):
    """Raised when an input is unsupported, malformed, or exceeds safety limits."""


@dataclass
class IngestedDocument:
    source_name: str
    source_type: str
    source_sha256: str
    text: str = ""
    rows: list[dict[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    extracted_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    confidence: float | None = None
    confidence: float | None = None


MAX_FILE_BYTES = 25 * 1024 * 1024
TEXT_EXTENSIONS = {".txt", ".text", ".md", ".log"}
CSV_EXTENSIONS = {".csv"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
PDF_EXTENSIONS = {".pdf"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | CSV_EXTENSIONS | XLSX_EXTENSIONS | PDF_EXTENSIONS


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise IngestionError("Text encoding could not be identified. Save as UTF-8 or provide a supported file.")


def _safe_name(name: str) -> str:
    return Path(name.replace("\\", "/")).name or "pasted-input.txt"


def ingest_bytes(name: str, data: bytes, *, pasted_text: bool = False) -> IngestedDocument:
    """Parse supported Phase 1 inputs without changing the original source bytes."""
    if not isinstance(data, bytes):
        raise IngestionError("Input must be bytes.")
    if not data:
        raise IngestionError("Input is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise IngestionError(f"Input exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB limit.")

    safe_name = _safe_name(name)
    suffix = Path(safe_name).suffix.lower()
    digest = hashlib.sha256(data).hexdigest()

    if pasted_text:
        text = _decode_text(data)
        if not text.strip():
            raise IngestionError("Pasted text is blank.")
        return IngestedDocument(safe_name, "pasted_text", digest, text=text)

    if suffix in TEXT_EXTENSIONS:
        text = _decode_text(data)
        return IngestedDocument(safe_name, "text", digest, text=text)
    if suffix in CSV_EXTENSIONS:
        try:
            decoded = _decode_text(data)
            sample = decoded[:8192]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
            if not reader.fieldnames:
                raise IngestionError("CSV must contain a header row.")
            rows = []
            for row in reader:
                rows.append({str(k).strip(): (v or "").strip() for k, v in row.items() if k is not None})
            return IngestedDocument(safe_name, "csv", digest, text=decoded, rows=rows)
        except csv.Error as exc:
            raise IngestionError(f"CSV parsing failed: {exc}") from exc
    if suffix in XLSX_EXTENSIONS:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise IngestionError("Excel support requires openpyxl. Install the FDIA dependencies.") from exc
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows: list[dict[str, str]] = []
            warnings: list[str] = []
            for sheet in workbook.worksheets:
                iterator = sheet.iter_rows(values_only=True)
                headers = next(iterator, None)
                if not headers:
                    continue
                header_names = [str(value).strip() if value is not None else "" for value in headers]
                if not any(header_names):
                    warnings.append(f"Sheet '{sheet.title}' has no usable header row.")
                    continue
                for row_number, values in enumerate(iterator, start=2):
                    if not any(value is not None for value in values):
                        continue
                    item = {"_sheet": sheet.title, "_row": str(row_number)}
                    for index, header in enumerate(header_names):
                        if header:
                            value = values[index] if index < len(values) else None
                            item[header] = "" if value is None else str(value).strip()
                    rows.append(item)
            workbook.close()
            return IngestedDocument(safe_name, "xlsx", digest, rows=rows, warnings=warnings)
        except IngestionError:
            raise
        except Exception as exc:
            raise IngestionError(f"Excel workbook could not be parsed: {exc}") from exc
    if suffix in PDF_EXTENSIONS:
        if not data.startswith(b"%PDF-"):
            raise IngestionError("File extension says PDF, but the file signature is invalid.")
        try:
            from pypdf import PdfReader
        except ImportError:
            return IngestedDocument(
                safe_name, "pdf", digest,
                warnings=["PDF signature accepted, but text extraction is unavailable: install pypdf. Scanned PDFs require the OCR phase."],
            )
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            page_text = []
            for index, page in enumerate(reader.pages, start=1):
                page_text.append(f"[Page {index}]\\n{page.extract_text() or ''}")
            text = "\\n\\n".join(page_text).strip()
            warnings = [] if text else ["No embedded text found. This may be a scanned PDF and needs OCR."]
            return IngestedDocument(safe_name, "pdf", digest, text=text, warnings=warnings)
        except Exception as exc:
            raise IngestionError(f"PDF parsing failed: {exc}") from exc

    raise IngestionError(
        f"Unsupported file type '{suffix or '[no extension]'}'. Phase 1 supports TXT, MD, LOG, CSV, XLSX, XLSM and text-based PDF."
    )


def normalize_header(value: str) -> str:
    """Normalize source labels for matching while retaining original labels in evidence."""
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def validate_record(record: dict[str, Any], *, required_fields: tuple[str, ...] = ()) -> dict[str, Any]:
    """Return deterministic validation findings; never modifies or submits the record."""
    errors: list[str] = []
    warnings: list[str] = []
    for field_name in required_fields:
        value = record.get(field_name)
        if value is None or not str(value).strip():
            errors.append(f"Missing required field: {field_name}")

    for key in ("imei", "imei1", "imei2"):
        value = record.get(key)
        if value is not None and str(value).strip():
            digits = re.sub(r"\s+", "", str(value))
            if not digits.isdigit() or len(digits) != 15:
                errors.append(f"{key} must contain exactly 15 digits.")

    phone = record.get("phone") or record.get("mobile")
    if phone is not None and str(phone).strip():
        digits = re.sub(r"\D", "", str(phone))
        if len(digits) not in (10, 11, 12, 13, 14, 15):
            errors.append("Phone/mobile number must contain 10–15 digits after removing formatting.")

    for key in ("amount", "rate", "quantity", "taxable_value", "gst_amount", "total"):
        value = record.get(key)
        if value is None or str(value).strip() == "":
            continue
        try:
            number = Decimal(str(value).replace(",", "").replace("₹", "").strip())
            if not number.is_finite():
                raise InvalidOperation
            if key in {"amount", "rate", "quantity", "taxable_value", "gst_amount", "total"} and number < 0:
                errors.append(f"{key} cannot be negative.")
        except (InvalidOperation, ValueError):
            errors.append(f"{key} must be a valid numeric value.")

    return {"valid": not errors, "errors": errors, "warnings": warnings}


def dry_run_preview(records: list[dict[str, Any]], *, target_profile: str, module: str) -> dict[str, Any]:
    """Build a safe preview payload. This function never performs network or ERP writes."""
    items = []
    for index, record in enumerate(records, start=1):
        validation = validate_record(record)
        items.append({
            "row": index,
            "operation": "create_or_update_proposed",
            "target_profile": target_profile,
            "module": module,
            "status": "READY_FOR_REVIEW" if validation["valid"] else "BLOCKED",
            "fields": record,
            "validation": validation,
            "writes_performed": False,
        })
    return {
        "mode": "DRY_RUN",
        "target_profile": target_profile,
        "module": module,
        "record_count": len(items),
        "writes_performed": False,
        "items": items,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def export_reconciliation_json(preview: dict[str, Any]) -> str:
    """Serialize a preview/outcome report; sensitive data should be redacted by the caller."""
    return json.dumps(preview, ensure_ascii=False, indent=2, default=str)
